#!/usr/bin/env python3
"""Build src/data.generated.js and data/VALIDATION.md from the raw tables.

Inputs (data/raw/, downloaded as published, CC BY 4.0):
  abs-sih-2019-20-2.xlsx   ABS Household Income and Wealth 2019-20, cube 2
  abs-sih-2019-20-7.xlsx   ... cube 7 (net worth quintiles)
  abs-sih-2019-20-10.xlsx  ... cube 10 (age of reference person)
  pc-wealth-transfers-chart-data.xlsx  Productivity Commission (2021),
                           Wealth transfers and their economic effects
  abs-total-value-dwellings-643201.xlsx  ABS Total Value of Dwellings (6432.0)
  oecd-house-prices-aus.csv  OECD Analytical house prices indicators, Australia
  rba-f5-lending-rates.csv   RBA statistical table F5, indicator lending rates
  rba-g1-cpi.csv             RBA statistical table G1, consumer price inflation

Run: pip install -r scripts/requirements.txt && python3 scripts/build_data.py
"""

import csv
import datetime
import json
import math
import os

import numpy as np
import openpyxl
from scipy.optimize import minimize
from scipy.stats import norm

ROOT = os.path.dirname(os.path.dirname(os.path.abspath(__file__)))
RAW = os.path.join(ROOT, 'data', 'raw')


def sheet(name, tab):
    return openpyxl.load_workbook(os.path.join(RAW, name), read_only=True, data_only=True)[tab]


def rows(ws, max_row=200):
    return [list(r) for r in ws.iter_rows(max_row=max_row, values_only=True)]


def find(rs, label, start=0):
    """First row at or after `start` whose first text cell starts with `label`."""
    for i in range(start, len(rs)):
        text = next((c for c in rs[i] if isinstance(c, str)), None)
        if text and text.strip().startswith(label):
            return i
    raise KeyError(label)


def nums(row):
    return [c for c in row if isinstance(c, (int, float))]


# ---------------------------------------------------------------- ABS ----

AGE_BANDS = ['15–24', '25–34', '35–44', '45–54', '55–64', '65–74', '75+']
# Where each band sits on the animation's age axis (20-90).
AGE_MID = [21.5, 29.5, 39.5, 49.5, 59.5, 69.5, 82]

t102 = rows(sheet('abs-sih-2019-20-10.xlsx', 'Table 10.2'), 60)
age_mean = nums(t102[find(t102, 'Mean net worth')])[:7]
age_median = nums(t102[find(t102, 'Median net worth')])[:7]
t103 = rows(sheet('abs-sih-2019-20-10.xlsx', 'Table 10.3'), 75)
households = nums(t103[find(t103, 'Households', find(t103, 'Estimated number'))])
age_households, total_households = households[:7], households[8]

t21 = rows(sheet('abs-sih-2019-20-2.xlsx', 'Table 2.1'), 45)
quintile_mean = [nums(t21[find(t21, l)])[-1] for l in
                 ['Lowest quintile', 'Second quintile', 'Third quintile',
                  'Fourth quintile', 'Highest quintile']]
pct = [nums(t21[find(t21, f'{10 * k}th (P{10 * k})')])[-1] for k in range(1, 10)]
all_mean = nums(t21[find(t21, 'All households')])[-1]

t72 = rows(sheet('abs-sih-2019-20-7.xlsx', 'Table 7.2'), 40)
accounts = nums(t72[find(t72, 'Value of accounts held with financial institutions')])[:5]
offsets = nums(t72[find(t72, 'Value of offset accounts')])[:5]
cash = [a + o for a, o in zip(accounts, offsets)]  # $'000

# The national quantile function, through ABS's published P10..P90.
# Below P20: linear in dollars (net worth crosses zero), with P0 set so the
# lowest quintile's mean is ABS's. P20..P90: log-linear. Above P90: a Pareto
# tail whose mean makes the highest quintile's mean ABS's.
P0 = 4 * quintile_mean[0] - 2 * pct[0] - pct[1]
d9_mean = (pct[8] - pct[7]) / math.log(pct[8] / pct[7])
d10_mean = 2 * quintile_mean[4] - d9_mean
alpha = d10_mean / (d10_mean - pct[8])
knots = [P0] + pct


def q_abs(p):
    p = min(max(p, 0.0), 1 - 1e-9)
    if p >= 0.9:
        return pct[8] * (0.1 / (1 - p)) ** (1 / alpha)
    k = int(p * 10)
    u = p * 10 - k
    a, b = knots[k], knots[k + 1]
    if k < 2:
        return a + (b - a) * u
    return a * (b / a) ** u


# Each age band: a shifted lognormal, X = exp(mu + s Z) - c. For any shift c
# the band's ABS mean and median pin mu and s exactly; the seven shifts are
# then fitted so the bands, mixed by household count, come closest to ABS's
# national percentiles and quintile means.
SAMPLES = 4000
Z = norm.ppf((np.arange(SAMPLES) + 0.5) / SAMPLES)
W = np.array(age_households) / sum(age_households)
PCT, QM = np.array(pct), np.array(quintile_mean)


def band_samples(c):
    out = []
    for m, md, cb in zip(age_mean, age_median, c):
        mu = math.log(md + cb)
        s = math.sqrt(2 * math.log((m + cb) / (md + cb)))
        out.append(np.exp(mu + s * Z) - cb)
    return out


def pooled(bs):
    xs = np.concatenate(bs)
    ws = np.concatenate([np.full(SAMPLES, w / SAMPLES) for w in W])
    o = np.argsort(xs, kind='stable')
    return xs[o], ws[o], o


def national(bs):
    xs, ws, _ = pooled(bs)
    cw = np.cumsum(ws)
    p = np.interp(np.arange(1, 10) / 10, cw, xs)
    q = []
    for k in range(5):
        sel = (cw > k / 5) & (cw <= (k + 1) / 5)
        q.append((xs[sel] * ws[sel]).sum() / ws[sel].sum())
    return p, np.array(q)


def loss(log_c):
    p, q = national(band_samples(np.exp(log_c)))
    return (np.sum(((p - PCT) / np.maximum(PCT, 100)) ** 2)
            + np.sum(((q - QM) / np.maximum(QM, 100)) ** 2))


fit = minimize(loss, np.log(np.full(7, 50.0)), method='Nelder-Mead',
               options={'maxiter': 4000, 'xatol': 1e-4, 'fatol': 1e-10})
shifts = np.exp(fit.x)
fit_bands = band_samples(shifts)
fit_pct, fit_qm = national(fit_bands)

# Then move every household to the national quantile of its rank in the
# pooled fit: the national distribution becomes ABS's exactly, and since the
# fit is already close, the bands move little.
xs, ws, order = pooled(fit_bands)
ranks = np.cumsum(ws) - ws / 2
mapped = np.empty_like(xs)
mapped[order] = [q_abs(r) for r in ranks]
bands = [mapped[b * SAMPLES:(b + 1) * SAMPLES] for b in range(7)]
everyone = sorted((float(v), float(W[b] / SAMPLES)) for b in range(7) for v in bands[b])


def decile_means(xs):
    n = len(xs)
    xs = sorted(xs)
    return [sum(xs[k * n // 10:(k + 1) * n // 10]) / (n // 10) for k in range(10)]


band_deciles = [decile_means(b) for b in bands]

# ---------------------------------------------------------------- PC -----

PCX = 'pc-wealth-transfers-chart-data.xlsx'


def table(tab, header, ncols, max_row=80):
    """Rows under the row containing `header`, first `ncols` non-empty cells."""
    rs = rows(sheet(PCX, tab), max_row)
    i = next(i for i, r in enumerate(rs) if header in [c for c in r if isinstance(c, str)])
    out = []
    for r in rs[i + 1:]:
        cells = [c for c in r if c is not None]
        if len(cells) < ncols:
            break
        out.append(cells[:ncols])
    return out


fig3 = table('Figure 3', 'Age band', 3)  # age band, % of gifts, % of inheritances
fig14 = table('1.4', 'Average inheritance', 3)  # age, average, %
fig110 = table('1.10', 'Average gift', 3)
transfer_ages = [r[0] for r in fig3]
gift_share = [r[1] / 100 for r in fig3]
inh_share = [r[2] / 100 for r in fig3]
avg_inh = {r[0]: r[1] for r in fig14}
avg_gift = {r[0]: r[1] for r in fig110}
# Gifts past 54 are "negligible" and unreported; carry the last average on.
last = None
gift_avg_by_age = []
for a in transfer_ages:
    last = avg_gift.get(a, last)
    gift_avg_by_age.append(last)
inh_avg_by_age = [avg_inh[a] for a in transfer_ages]


def pc_rows(tab, max_row=80):
    return rows(sheet(PCX, tab), max_row)


def quintile_block(rs, col):
    """Five quintile values in column `col` under a 'Quintile' header."""
    i = next(i for i, r in enumerate(rs) if 'Quintile' in r)
    return [rs[i + 1 + k][col] for k in range(5)]


r21 = pc_rows('2.1')
qcols = [j for j, c in enumerate(next(r for r in r21 if 'Quintile' in r)) if c == 'Quintile']
inh_recv_share = [v / 100 for v in quintile_block(r21, qcols[0] + 1)]
inh_recv_avg = quintile_block(r21, qcols[1] + 1)
r23 = pc_rows('2.3')
qcols = [j for j, c in enumerate(next(r for r in r23 if 'Quintile' in r)) if c == 'Quintile']
gift_all_avg = quintile_block(r23, qcols[0] + 1)
gift_recv_share = [v / 100 for v in quintile_block(r23, qcols[1] + 1)]
gift_recv_avg = quintile_block(r23, qcols[2] + 1)
r4 = pc_rows('Figure 4')
qcols = [j for j, c in enumerate(next(r for r in r4 if 'Quintile' in r)) if c == 'Quintile']
transfer_all_avg = quintile_block(r4, qcols[0] + 1)
transfer_pct_wealth = quintile_block(r4, qcols[1] + 1)

r13 = pc_rows('1.3', 80)
i = next(i for i, r in enumerate(r13) if '2018-19' in r)
row = r13[i]
y = [j for j, c in enumerate(row) if c == '2018-19']
inh_value_1819 = row[y[0] + 2]
inh_count_1819 = row[y[1] + 1]
inh_avg_1819 = row[y[1] + 2]
r19 = pc_rows('1.9', 80)
i = next(i for i, r in enumerate(r19) if '2018-19' in r)
row = r19[i]
y = [j for j, c in enumerate(row) if c == '2018-19']
gift_value_1819 = row[y[0] + 3]
gift_parent_value_1819 = row[y[0] + 1]
gift_count_1819 = row[y[1] + 1]
gift_avg_1819 = row[y[1] + 2]

# Estimated inheritance outflow to the next generation (figure 1.7), 2018:
# the PC's HILDA-based estimate, and Roach (2018) as scaled by the PC (which
# notes the scaling "will tend to inflate its value"). The 2018 row holds the
# HILDA value; the Roach value sits on the row after it, in its own column.
r17 = pc_rows('1.7', 60)
i = next(i for i, r in enumerate(r17) if 'Administrative estimate' in r)
hdr = r17[i]
c_year, c_hilda, c_roach = hdr.index('Year'), hdr.index('HILDA estimate'), hdr.index('Roach (2018)')
k = next(k for k in range(i + 1, len(r17)) if r17[k][c_year] == 2018)
outflow_2018 = {'hilda': r17[k][c_hilda], 'roach': r17[k + 1][c_roach]}
assert isinstance(outflow_2018['hilda'], float) and isinstance(outflow_2018['roach'], float), outflow_2018

r210 = pc_rows('2.10', 80)
i = next(i for i, r in enumerate(r210) if 'Age of child' in r)
c = r210[i].index('Age of child')
persistence = [[r[c], r[c + 1]] for r in r210[i + 1:] if isinstance(r[c], (int, float)) and r[c + 1] is not None]

# ----------------------------------------------------------- housing -----

WEEKS = 365.25 / 7

# Household disposable income by decile within each age band. ABS table 10.1
# gives equivalised disposable income at P10..P90 by age; a decile's mean is
# taken at its middle (P15 = halfway between P10 and P20, and so on), the
# bottom decile's at P5 extrapolated from P10 and P20, and the top decile's
# so the ten average to the band's published mean. Equivalised is turned
# back into household income by the band's mean disposable ÷ mean
# equivalised disposable.
t101 = rows(sheet('abs-sih-2019-20-10.xlsx', 'Table 10.1'), 45)
mean_disp = nums(t101[find(t101, 'Disposable income')])[:7]
mean_eq = nums(t101[find(t101, 'Equivalised disposable income')])[:7]
housing_costs = nums(t101[find(t101, 'Housing costs')])[:7]
median_disp_all = nums(t101[find(t101, 'Disposable income', find(t101, 'Median income'))])[8]
eq_pct = [nums(t101[find(t101, f'{10 * k}th (P{10 * k})')])[:7] for k in range(1, 10)]
income_deciles = []
for b in range(7):
    p = [eq_pct[k][b] for k in range(9)]
    d = [p[0] - (p[1] - p[0]) / 2] + [(p[k] + p[k + 1]) / 2 for k in range(8)]
    d.append(10 * mean_eq[b] - sum(d))
    scale = mean_disp[b] / mean_eq[b]
    income_deciles.append([round(v * scale * WEEKS) for v in d])

# Tenure by age (validation): owners outright and with a mortgage.
outright_by_age = nums(t103[find(t103, 'Owner without a mortgage')])[:7]
mortgaged_by_age = nums(t103[find(t103, 'Owner with a mortgage')])[:7]

# House price: ABS mean price of residential dwellings, Australia, averaged
# over the four quarters of 2019-20 (ABS publishes a mean, not a median).
ws = sheet('abs-total-value-dwellings-643201.xlsx', 'Data1')
head = next(ws.iter_rows(max_row=1, values_only=True))
col = next(j for j, c in enumerate(head) if c and str(c).split(';')[0].strip() == 'Mean price of residential dwellings'
           and str(c).split(';')[1].strip() == 'Australia')
quarters = [r[col] for r in ws.iter_rows(min_row=11, values_only=True)
            if isinstance(r[0], datetime.datetime)
            and datetime.datetime(2019, 7, 1) <= r[0] <= datetime.datetime(2020, 6, 30)]
assert len(quarters) == 4, quarters
price_2019_20 = sum(quarters) / 4 * 1e3

# Price-to-income ratio by year (OECD, index 2015 = 100), relative to
# 2019-20 (the mean of 2019 and 2020).
ratio = {}
with open(os.path.join(RAW, 'oecd-house-prices-aus.csv')) as f:
    for r in csv.DictReader(f):
        if r['MEASURE'] == 'HPI_YDH':
            ratio[int(r['TIME_PERIOD'])] = float(r['OBS_VALUE'])
ratio_base = (ratio[2019] + ratio[2020]) / 2
price_by_year = {y: round(price_2019_20 * v / ratio_base) for y, v in sorted(ratio.items())}

# Mortgage rate: the RBA's standard variable owner-occupier rate, averaged
# over every month in the OECD series' years.
with open(os.path.join(RAW, 'rba-f5-lending-rates.csv'), encoding='utf-8-sig') as f:
    rba = list(csv.reader(f))
rcol = rba[1].index('Lending rates; Housing loans; Banks; Variable; Standard; Owner-occupier')
dcol = rba[1].index('Lending rates; Housing loans; Banks; Variable; Discounted; Owner-occupier')
monthly = [(int(r[0][-4:]), float(r[rcol])) for r in rba if r and r[0][:2].isdigit() and r[rcol]]
y0, y1 = min(ratio), max(ratio)
span = [v for y, v in monthly if y0 <= y <= y1]
rate_long_run = sum(span) / len(span) / 100
# Incomes and prices here are 2019-20 dollars, so the model needs the real
# rate: each year's average nominal rate deflated by that year's CPI growth
# (annual average CPI over the previous year's), averaged over the span.
with open(os.path.join(RAW, 'rba-g1-cpi.csv'), encoding='utf-8-sig') as f:
    g1 = list(csv.reader(f))
cpi = {}
for r in g1:
    if len(r) > 1 and r[0][:2].isdigit() and r[1]:
        cpi.setdefault(int(r[0][-4:]), []).append(float(r[1]))
cpi = {y: sum(v) / len(v) for y, v in cpi.items()}
# Rates actually paid: the discounted owner-occupier rate where the RBA has
# it (from 2004), the standard rate before.
cell = lambda r, c: r[c] if c < len(r) else ''
paid = [(int(r[0][-4:]), float(cell(r, dcol) or cell(r, rcol))) for r in rba
        if r and r[0][:2].isdigit() and (cell(r, dcol) or cell(r, rcol))]
nominal = {}
for y, v in paid:
    nominal.setdefault(y, []).append(v / 100)
nominal = {y: sum(v) / len(v) for y, v in nominal.items()}
first_discounted = min(int(r[0][-4:]) for r in rba if r and r[0][:2].isdigit() and cell(r, dcol))
real_by_year = {y: (1 + nominal[y]) / (cpi[y] / cpi[y - 1]) - 1
                for y in range(y0, y1 + 1) if y in nominal and y in cpi and y - 1 in cpi}
real_rate = sum(real_by_year.values()) / len(real_by_year)
paid_nominal = sum(nominal[y] for y in real_by_year) / len(real_by_year)
real_early = [v for y, v in real_by_year.items() if y < 1990]
real_late = [v for y, v in real_by_year.items() if y >= 1990]
inflation = sum(cpi[y] / cpi[y - 1] - 1 for y in real_by_year) / len(real_by_year)
rate_by_year = {}
for y, v in monthly:
    rate_by_year.setdefault(y, []).append(v)
rate_by_year = {y: round(sum(v) / len(v) / 100, 4) for y, v in rate_by_year.items() if y0 <= y <= y1}

# ------------------------------------------------------------ output -----

data = {
    'ageBands': AGE_BANDS,
    'ageMid': AGE_MID,
    'ageMean': [v * 1e3 for v in age_mean],
    'ageMedian': [v * 1e3 for v in age_median],
    'ageHouseholds': [v * 1e3 for v in age_households],
    'households': total_households * 1e3,
    'bandDeciles': [[round(v * 1e3) for v in d] for d in band_deciles],
    'nationalPercentiles': [v * 1e3 for v in pct],
    'quintileMean': [v * 1e3 for v in quintile_mean],
    'cashByQuintile': [round(v * 1e3) for v in cash],
    'transferAges': transfer_ages,
    'giftShareByAge': gift_share,
    'inheritanceShareByAge': inh_share,
    'giftAvgByAge': [round(v) for v in gift_avg_by_age],
    'inheritanceAvgByAge': [round(v) for v in inh_avg_by_age],
    'inheritanceRecvShareByQuintile': inh_recv_share,
    'inheritanceRecvAvgByQuintile': [round(v) for v in inh_recv_avg],
    'giftRecvShareByQuintile': gift_recv_share,
    'giftRecvAvgByQuintile': [round(v) for v in gift_recv_avg],
    'inheritancesPerYear': round(inh_count_1819),
    'giftsPerYear': round(gift_count_1819),
    'inheritanceValue': round(inh_value_1819),
    'giftValue': round(gift_value_1819),
    'giftFromParentsValue': round(gift_parent_value_1819),
    'inheritanceOutflow2018': {k: round(v) for k, v in outflow_2018.items()},
    'persistenceByAge': persistence,
    'incomeDeciles': income_deciles,
    'medianDisposable': round(median_disp_all * WEEKS),
    'housingCostShareByAge': [round(h / d, 4) for h, d in zip(housing_costs, mean_disp)],
    'outrightByAge': [v / 100 for v in outright_by_age],
    'mortgagedByAge': [v / 100 for v in mortgaged_by_age],
    'price2019_20': round(price_2019_20),
    'priceByYear': price_by_year,
    'rateLongRun': round(rate_long_run, 4),
    'realRate': round(real_rate, 4),
    'paidNominalRate': round(paid_nominal, 4),
    'rateYears': [y0, y1],
    'rateByYear': rate_by_year,
}

with open(os.path.join(ROOT, 'src', 'data.generated.js'), 'w') as f:
    f.write('// GENERATED by scripts/build_data.py from data/raw/. Do not edit.\n')
    f.write('// Sources and method: data/SOURCES.md; checks: data/VALIDATION.md.\n')
    f.write('export const DATA = ')
    f.write(json.dumps(data, indent=1))
    f.write(';\n')

# ------------------------------------------------------------ checks -----


def k(x):
    return f'${x:,.0f}k'


def pct_err(model, real):
    return f'{100 * (model - real) / real:+.1f}%'


everyone.sort()
total_w = sum(w for _, w in everyone)


def quantile_of(sorted_w, p):
    acc = 0
    for v, w in sorted_w:
        acc += w
        if acc >= p * total_w:
            return v
    return sorted_w[-1][0]


def weighted_slice_mean(sorted_w, p0, p1):
    acc = s = ws = 0
    for v, w in sorted_w:
        lo, hi = acc, acc + w
        acc = hi
        take = max(0, min(hi, p1 * total_w) - max(lo, p0 * total_w))
        s += v * take
        ws += take
    return s / ws


lines = ['# Validation', '',
         'Generated by `scripts/build_data.py`. The model is checked against every',
         'ABS figure it was built from.', '',
         '## Stage 1: a shifted lognormal per age band', '',
         'Each band matches its ABS mean and median exactly; the shift c lets net',
         'worth go below zero. The shifts are fitted to the national figures.', '',
         '| Age band | ABS mean | ABS median | Households | Shift c |', '|---|---|---|---|---|']
for b, m, md, h, c in zip(AGE_BANDS, age_mean, age_median, age_households, shifts):
    lines.append(f'| {b} | {k(m)} | {k(md)} | {h:,.0f}k | {k(c)} |')
lines += ['', '| | Stage 1 | ABS 2019-20 | Error |', '|---|---|---|---|']
for i, real in enumerate(pct):
    lines.append(f'| P{10 * (i + 1)} | {k(fit_pct[i])} | {k(real)} | {pct_err(fit_pct[i], real)} |')
for i, real in enumerate(quintile_mean):
    lines.append(f'| Quintile {i + 1} mean | {k(fit_qm[i])} | {k(real)} | {pct_err(fit_qm[i], real)} |')
lines += ['', '## Stage 2: ranks mapped onto ABS\'s national percentiles', '',
          'Below P20 the national quantile function is linear in dollars (net worth',
          'crosses zero), P20–P90 log-linear, above P90 a Pareto tail; P0 and the tail',
          'are set so the lowest and highest quintile means are ABS\'s.', '',
          '| | Model | ABS 2019-20 | Error |', '|---|---|---|---|']
for i, real in enumerate(pct):
    m = quantile_of(everyone, (i + 1) / 10)
    lines.append(f'| P{10 * (i + 1)} | {k(m)} | {k(real)} | {pct_err(m, real)} |')
for i, real in enumerate(quintile_mean):
    m = weighted_slice_mean(everyone, i / 5, (i + 1) / 5)
    lines.append(f'| Quintile {i + 1} mean | {k(m)} | {k(real)} | {pct_err(m, real)} |')
m_all = weighted_slice_mean(everyone, 0, 1)
lines.append(f'| Mean | {k(m_all)} | {k(all_mean)} | {pct_err(m_all, all_mean)} |')
lines += ['', f'Derived: P0 = {k(P0)}; top decile mean {k(d10_mean)}; Pareto α above P90 = {alpha:.2f}.',
          'Cross-check (not used): ACOSS & UNSW report the top 10% averaging $5.2M in 2022-23,',
          'three years of price growth after this survey.', '',
          '## Age bands after stage 2', '',
          '| Age band | Model mean | ABS mean | Error | Model median | ABS median | Error |',
          '|---|---|---|---|---|---|---|']
for b, xs, m, md in zip(AGE_BANDS, bands, age_mean, age_median):
    mm = sum(xs) / len(xs)
    mmd = sorted(xs)[len(xs) // 2]
    lines.append(f'| {b} | {k(mm)} | {k(m)} | {pct_err(mm, m)} | {k(mmd)} | {k(md)} | {pct_err(mmd, md)} |')
lines += ['', '## Decile means by age band (the rows of the animation)', '',
          '| Decile | ' + ' | '.join(AGE_BANDS) + ' |', '|---' * 8 + '|']
for d in range(10):
    lines.append(f'| {d + 1} | ' + ' | '.join(k(bd[d]) for bd in band_deciles) + ' |')

# Cross-checks against ACOSS & UNSW (2024), which the model was not fitted
# to (their figures are 2022-23, three years on from this survey).


def pooled_band(ix):
    vals = [(float(v), float(W[b] / SAMPLES)) for b in ix for v in bands[b]]
    return sorted(vals)


def share_mean(vals, p0, p1):
    tot = sum(w for _, w in vals)
    acc = s_ = ws_ = 0
    for v, w in vals:
        lo, hi = acc, acc + w
        acc = hi
        take = max(0, min(hi, p1 * tot) - max(lo, p0 * tot))
        s_ += v * take
        ws_ += take
    return s_ / ws_


young = pooled_band([0, 1])
old = pooled_band([5, 6])
old_total = sum(v * w for v, w in old)
top6 = share_mean(old, 5 / 6, 1) * sum(w for _, w in old) / 6
lines += ['', '## Cross-checks (not fitted)', '',
          '| | Model (2019-20) | ACOSS & UNSW (2022-23) |', '|---|---|---|',
          f'| Under 35: lowest 60% mean | {k(share_mean(young, 0, 0.6))} | $80k |',
          f'| Under 35: top 10% mean | {k(share_mean(young, 0.9, 1))} | $2,000k |',
          f'| Over 65: share of their wealth held by the top sixth | {100 * top6 / old_total:.0f}% | 53% |',
          f'| Over 65: mean of that top sixth | {k(share_mean(old, 5 / 6, 1))} | $5,600k |']

# Transfers: does age × quintile reproduce the PC's averages?
exp_inh = sum(s * a for s, a in zip(inh_share, inh_avg_by_age))
exp_gift = sum(s * a for s, a in zip(gift_share, gift_avg_by_age))
lines += ['', '## Transfers', '',
          f'- Inheritances 2018-19: {inh_count_1819:,.0f} received, ${inh_value_1819 / 1e9:.1f}B, '
          f'average ${inh_avg_1819:,.0f}. Age shares × age averages give ${exp_inh:,.0f} '
          f'(an all-years average, 2018 dollars).',
          f'- Gifts from parents 2018-19: {gift_count_1819:,.0f}, average ${gift_avg_1819:,.0f}; all gifts '
          f'${gift_value_1819 / 1e9:.1f}B. Age shares × age averages give ${exp_gift:,.0f}.',
          f'- Per household per year (÷ {total_households * 1e3:,.0f} ABS households): '
          f'{inh_count_1819 / (total_households * 1e3):.3f} inheritances, '
          f'{gift_count_1819 / (total_households * 1e3):.3f} gifts — HILDA counts people, so this '
          'mixes people and households.',
          '- Average transfer received over three years, all people, by wealth quintile '
          '(PC figure 4a): ' + ', '.join(f'${v:,.0f}' for v in transfer_all_avg) +
          f' — top/bottom {transfer_all_avg[4] / transfer_all_avg[0]:.1f}×. As % of wealth: ' +
          ', '.join(f'{v:.1f}%' for v in transfer_pct_wealth) + '.',
          '- Cash on hand (accounts + offset, ABS table 7.2) by quintile: ' +
          ', '.join(f'${v * 1e3:,.0f}' for v in cash) + '.']
lines += ['', '## Housing inputs', '',
          f'- Mean dwelling price 2019-20: ${price_2019_20:,.0f} (quarters: ' +
          ', '.join(f'${q:,.1f}k' for q in quarters) + ').',
          f'- Mortgage rate, long-run average of the RBA standard variable rate {y0}–{y1}: '
          f'{100 * rate_long_run:.2f}% nominal (headline). Rates actually paid (the discounted rate from '
          f'{first_discounted}, the standard rate before): {100 * paid_nominal:.2f}% nominal; less each year\'s '
          f'CPI inflation (average {100 * inflation:.2f}%), a real rate of {100 * real_rate:.2f}%, which the model '
          f'uses (its money is 2019-20 dollars). Real rates by era: {100 * sum(real_early) / len(real_early):.1f}% '
          f'before 1990 (regulated, rationed credit), {100 * sum(real_late) / len(real_late):.1f}% from 1990.',
          '- House price at 2019-20 incomes for each year\'s price-to-income ratio: ' +
          ', '.join(f'{y} ${price_by_year[y] / 1e3:,.0f}k' for y in (1970, 1985, 1995, 2005, 2015, 2024)) + '.',
          '', '| Income decile | ' + ' | '.join(AGE_BANDS) + ' |', '|---' * 8 + '|']
for d in range(10):
    lines.append(f'| {d + 1} | ' + ' | '.join(f'${income_deciles[b][d] / 1e3:,.0f}k' for b in range(7)) + ' |')
lines.append('')
lines.append('Household disposable income a year; each band\'s ten deciles average to its ABS mean.')
with open(os.path.join(ROOT, 'data', 'VALIDATION.md'), 'w') as f:
    f.write('\n'.join(lines) + '\n')

print('\n'.join(lines))
